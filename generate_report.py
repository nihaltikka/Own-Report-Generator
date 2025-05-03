import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext
from datetime import datetime
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx.enum.style import WD_STYLE_TYPE
from reportlab.lib import colors
from fpdf import FPDF
import os
from datetime import datetime
import webbrowser

# Professional Blue Theme
COLORS = {
    'primary': '#2B579A',
    'secondary': '#3C78D8',
    'background': '#F0F4F8',
    'accent': '#4A90E2',
    'text': '#2D3748',
    'success': '#48BB78',
    'warning': '#ED8936',
    'error': '#F56565',
    'footer': '#718096'
}

class ModernPDF(FPDF):
    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=15)
        self.set_title("Vulnerability Report")
    
    def header(self):
        self.set_font('Arial', 'B', 16)
        self.set_text_color(*self.hex_to_rgb(COLORS['primary']))
        self.cell(0, 10, self.title, 0, 1, 'C')
        self.ln(10)
    
    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 10)
        self.set_text_color(*self.hex_to_rgb(COLORS['footer']))
        self.cell(0, 10, f'Report Prepared by XploitiX - Page {self.page_no()}', 0, 0, 'C')
    
    @staticmethod
    def hex_to_rgb(hex_color):
        hex_color = hex_color.lstrip('#')
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

class ReportGeneratorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("XploitiX Professional Reporter")
        self.vulnerabilities = []
        self.setup_ui()
        self.configure_styles()

    def configure_styles(self):
        style = ttk.Style()
        style.theme_use('alt')
        
        # Configure modern styles
        style.configure('.', background=COLORS['background'], foreground=COLORS['text'])
        style.configure('TFrame', background=COLORS['background'])
        style.configure('TLabel', font=('Segoe UI', 10), padding=5)
        style.configure('TEntry', fieldbackground='white', bordercolor=COLORS['secondary'], 
                      borderwidth=2, relief='flat', padding=5)
        style.configure('TCombobox', fieldbackground='white', bordercolor=COLORS['secondary'], 
                      arrowsize=15, padding=5)
        style.configure('TButton', font=('Segoe UI', 10, 'bold'), borderwidth=0, 
                      background=COLORS['secondary'], foreground='white', padding=10)
        style.map('TButton', background=[('active', COLORS['primary']), ('disabled', '#cccccc')])
    
    def setup_ui(self):
        self.main_frame = ttk.Frame(self.root)
        self.main_frame.pack(padx=25, pady=25, fill=tk.BOTH, expand=True)

        # Header Section
        header_frame = ttk.Frame(self.main_frame)
        header_frame.grid(row=0, column=0, columnspan=2, pady=(0, 20))
        
        ttk.Label(header_frame, text="XploitiX", font=('Segoe UI', 24, 'bold'), 
                foreground=COLORS['primary']).pack(side=tk.LEFT)
        ttk.Label(header_frame, text="Vulnerability Report Generator", font=('Segoe UI', 16), 
                foreground=COLORS['text']).pack(side=tk.LEFT, padx=10)

        # Report Name Entry (Fixed attribute name)
        info_frame = ttk.Frame(self.main_frame)
        info_frame.grid(row=1, column=0, columnspan=2, pady=10, sticky='ew')
        
        ttk.Label(info_frame, text="Report Name:", font=('Segoe UI', 11)).pack(side=tk.LEFT, padx=5)
        self.report_name = ttk.Entry(info_frame, width=40, font=('Segoe UI', 11))
        self.report_name.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=5)
                
        # Vulnerability Form with Enhanced Dropdown
        form_frame = ttk.Frame(self.main_frame)
        form_frame.grid(row=2, column=0, columnspan=2, sticky='ew')
        
        fields = [
            ("Vulnerability Name", 1),
            ("Vulnerability Description", 4),
            ("Impact", 4),
            ("Affected URL", 1),
            ("Reference", 1),
            ("Proof of Concept", 4),
            ("Severity", 1),
            ("Recommendation", 4)
        ]
        
        self.entries = {}
        for idx, (field, height) in enumerate(fields):
            row_frame = ttk.Frame(form_frame)
            row_frame.grid(row=idx, column=0, sticky='ew', pady=3)
            
            ttk.Label(row_frame, text=f"{field}:", width=20, anchor=tk.W).pack(side=tk.LEFT)
            
            if height > 1:
                entry = scrolledtext.ScrolledText(row_frame, height=height, width=60,
                                                font=('Segoe UI', 10), wrap=tk.WORD,
                                                relief='solid', borderwidth=1)
            else:
                if field == "Severity":
                    entry = ttk.Combobox(row_frame, 
                                       values=["Low", "High", "Medium", "Critical"],  # Updated order
                                       width=58, 
                                       font=('Segoe UI', 10),
                                       state='readonly')
                    entry.current(0)
                else:
                    entry = ttk.Entry(row_frame, width=60, font=('Segoe UI', 10))
            
            entry.pack(side=tk.LEFT, expand=True, fill=tk.X)
            self.entries[field] = entry

        # Control Panel with Modern Buttons
        control_frame = ttk.Frame(self.main_frame)
        control_frame.grid(row=3, column=0, columnspan=2, pady=20, sticky='ew')
        
        ttk.Button(control_frame, 
                 text="➕ Add Vulnerability", 
                 command=self.add_vulnerability,
                 style='Accent.TButton').pack(side=tk.LEFT, padx=5)
                 
        ttk.Label(control_frame, text="Output Format:").pack(side=tk.LEFT, padx=20)
        self.format_var = tk.StringVar(value="PDF")
        formats = ttk.Combobox(control_frame, 
                             textvariable=self.format_var, 
                             values=["PDF", "Word", "Excel", "HTML"], 
                             width=15,
                             state='readonly')
        formats.pack(side=tk.LEFT)
        
        generate_btn = ttk.Button(control_frame, 
                                text="🛠 Generate Report", 
                                command=self.generate_report)  # This must match method name
        generate_btn.pack(side=tk.RIGHT)

        # Configure grid weights
        self.main_frame.columnconfigure(0, weight=1)
        form_frame.columnconfigure(0, weight=1)

    def add_vulnerability(self):
        data = self.process_inputs()
        if not self.validate_inputs(data):
            return
        self.vulnerabilities.append(data)
        self.clear_fields()
        messagebox.showinfo("Added", f"Vulnerability '{data['Vulnerability Name']}' added!\nTotal: {len(self.vulnerabilities)}")

    def process_inputs(self):
        data = {}
        for field, entry in self.entries.items():
            # Get value with proper casing for severity
            if field == "Severity":
                value = entry.get().strip().title()  # Force title case
            else:
                value = entry.get("1.0", tk.END).strip() if isinstance(entry, tk.Text) else entry.get()
            data[field] = value if value else "N/A"
        return data

    def validate_inputs(self, data):
        if data["Vulnerability Name"] == "N/A":
            messagebox.showwarning("Input Error", "Vulnerability Name is required")
            return False
        return True

    def clear_fields(self):
        for entry in self.entries.values():
            if isinstance(entry, tk.Text):
                entry.delete("1.0", tk.END)
            else:
                entry.delete(0, tk.END)

    def generate_report(self):
        try:
            if not self.vulnerabilities:
                messagebox.showwarning("Input Error", "Add at least one vulnerability")
                return

            report_meta = {
                "name": self.report_name.get().strip() or "Unnamed_Report",
                "date": datetime.now().strftime("%Y-%m-%d_%H-%M-%S"),
                "author": "XploitiX"
            }

            format_handlers = {
                "pdf": self.generate_pdf,
                "word": self.generate_word,
                "excel": self.generate_excel,
                "html": self.generate_html
            }

            format_selected = self.format_var.get().lower()
            if format_selected not in format_handlers:
                messagebox.showerror("Error", "Invalid output format selected")
                return

            try:
                filename = format_handlers[format_selected](report_meta)
                if filename and os.path.exists(filename):
                    messagebox.showinfo(
                        "Success", 
                        f"Report generated successfully!\n"
                        f"Location: {os.path.abspath(filename)}"
                    )
                else:
                    messagebox.showerror("Error", "Failed to generate report file")
                    
            except Exception as e:
                messagebox.showerror("Generation Error", f"Failed to create report: {str(e)}")

        except Exception as main_error:
            messagebox.showerror("System Error", f"Application error: {str(main_error)}")
    def get_pdf_styles(self):  # Correct spelling
        styles = getSampleStyleSheet()
        styles['Title'].textColor = colors.HexColor(COLORS['primary'])
        styles['Heading1'].textColor = colors.HexColor(COLORS['text'])
        styles['Heading2'].textColor = colors.HexColor(COLORS['secondary'])
        styles['Normal'].textColor = colors.HexColor(COLORS['text'])
        styles.add(ParagraphStyle(
            name='Footer',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor(COLORS['footer'])
        ))
        return styles

    def generate_pdf(self, meta):
        try:
            filename = f"{meta['name']}_{meta['date']}.pdf"
            doc = SimpleDocTemplate(filename, pagesize=letter)
            styles = self.get_pdf_styles()
            story = []
            
            # Header
            story.append(Paragraph(meta['name'], styles['Title']))
            story.append(Spacer(1, 10))
            story.append(Paragraph(f"<font color={COLORS['footer']}>Generated: {meta['date']}</font>", styles['Normal']))
            story.append(Spacer(1, 20))
            
            # Content
            for vuln in self.vulnerabilities:
                story += self.create_vulnerability_section(vuln, styles)
                story.append(Spacer(1, 15))
            
            # Footer
            story.append(Spacer(1, 20))
            story.append(Paragraph(f"<font color={COLORS['footer']}>Report Prepared by {meta['author']}</font>", 
                                styles['Footer']))
            
            doc.build(story)
            return os.path.abspath(filename)
        except Exception as e:
            messagebox.showerror("PDF Error", f"PDF generation failed: {str(e)}")
            raise

    def create_vulnerability_section(self, vuln, styles):
        section = []
        section.append(Paragraph(vuln['Vulnerability Name'], styles['Heading1']))
        for key, value in vuln.items():
            if key == "Vulnerability Name":
                continue
            section.append(Paragraph(f"<b>{key}:</b>", styles['Heading2']))
            if key == "Severity":
                section.append(Paragraph(f"<font color={self.get_severity_color(value)}>{value}</font>", 
                                      styles['Normal']))
            else:
                section.append(Paragraph(value, styles['Normal']))
            section.append(Spacer(1, 8))
        return section

    def generate_word(self, meta):
        try:
            filename = f"{meta['name'].replace(' ', '_')}_{meta['date']}.docx"
            doc = Document()
            
            # Create custom styles
            self.create_word_styles(doc)
            
            # Main title (paragraph style)
            title = doc.add_paragraph(meta['name'], style='Title')
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # Date (paragraph style)
            date_para = doc.add_paragraph(f"Generated: {meta['date']}", style='Subtitle')
            date_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # Vulnerability content
            for vuln in self.vulnerabilities:
                doc.add_heading(vuln['Vulnerability Name'], level=1)
                
                for key, value in vuln.items():
                    if key == "Vulnerability Name":
                        continue
                    
                    # Add heading (paragraph style)
                    doc.add_heading(key, level=2)
                    
                    # Add content with potential character styling
                    para = doc.add_paragraph()
                    run = para.add_run(value)
                    
                    # Apply character style only to severity
                    if key == "Severity":
                        run.font.color.rgb = self.word_severity_color(value)
                        run.bold = True

            # Footer (paragraph style)
            footer = doc.sections[0].footer
            footer_para = footer.add_paragraph(f"Report Prepared by {meta['author']}")
            footer_para.style = doc.styles['Footer']
            
            doc.save(filename)
            return os.path.abspath(filename)
        except Exception as e:
            messagebox.showerror("Word Error", f"Word generation failed: {str(e)}")
            raise

    def create_word_styles(self, doc):
        styles = doc.styles
        
        # Create proper paragraph styles
        if 'Footer' not in styles:
            footer_style = styles.add_style('Footer', WD_STYLE_TYPE.PARAGRAPH)
            footer_style.font.color.rgb = RGBColor.from_string(COLORS['footer'][1:])
            footer_style.font.size = Pt(10)
            footer_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Modify existing title style
        title_style = styles['Title']
        title_style.font.color.rgb = RGBColor.from_string(COLORS['primary'][1:])
        title_style.font.size = Pt(18)
        title_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def generate_excel(self, meta):
        try:
            filename = f"{meta['name'].replace(' ', '_')}_{meta['date']}.xlsx"
            df = pd.DataFrame(self.vulnerabilities)
            
            with pd.ExcelWriter(filename, engine='xlsxwriter') as writer:
                df.to_excel(writer, index=False, sheet_name='Vulnerabilities')
                workbook = writer.book
                worksheet = writer.sheets['Vulnerabilities']
                
                header_format = workbook.add_format({
                    'bold': True,
                    'bg_color': COLORS['primary'],
                    'font_color': '#FFFFFF',
                    'border': 1
                })
                
                for col_num, value in enumerate(df.columns.values):
                    worksheet.write(0, col_num, value, header_format)
                
                severity_col = df.columns.get_loc("Severity")
                for row_num, severity in enumerate(df['Severity'], 1):
                    fmt = workbook.add_format({
                        'font_color': self.get_severity_color(severity, excel=True)[1:],
                        'bold': True
                    })
                    worksheet.write(row_num, severity_col, severity, fmt)
                
                worksheet.write(len(df)+2, 0, f"Report Prepared by {meta['author']}", 
                              workbook.add_format({'font_color': COLORS['footer']}))
                
                worksheet.set_column('A:H', 25)
            
            return os.path.abspath(filename)
        except Exception as e:
            messagebox.showerror("Excel Error", f"Excel generation failed: {str(e)}")
            raise

    def generate_html(self, meta):
            try:
                filename = f"{meta['name'].replace(' ', '_')}_{meta['date']}.html"
                html_content = f"""<!DOCTYPE html>
    <html>
    <head>
        <title>{meta['name']}</title>
        <style>
            body {{ 
                background-color: {COLORS['background']}; 
                color: {COLORS['text']};
                font-family: 'Segoe UI', sans-serif;
                line-height: 1.6;
                margin: 2em;
            }}
            .header {{ 
                text-align: center;
                border-bottom: 2px solid {COLORS['primary']};
                padding-bottom: 1em;
                margin-bottom: 2em;
            }}
            .vuln {{
                background: #FFFFFF;
                padding: 1.5em;
                border-radius: 8px;
                margin-bottom: 2em;
                box-shadow: 0 2px 4px rgba(0,0,0,0.1);
            }}
            h1 {{ color: {COLORS['primary']}; }}
            h2 {{ color: {COLORS['text']}; }}
            h3 {{ color: {COLORS['secondary']}; }}
            .footer {{
                text-align: center;
                color: {COLORS['footer']};
                margin-top: 3em;
                padding-top: 1em;
                border-top: 1px solid {COLORS['footer']};
            }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>{meta['name']}</h1>
            <p>Generated: {meta['date']}</p>
        </div>"""

                # Add vulnerabilities
                for vuln in self.vulnerabilities:
                    html_content += f'<div class="vuln"><h2>{vuln["Vulnerability Name"]}</h2>'
                    for key, value in vuln.items():
                        if key == "Vulnerability Name":
                            continue
                        color_style = f'style="color: {self.get_severity_color(value)};"' if key == "Severity" else ""
                        html_content += f'<h3>{key}</h3><p {color_style}>{value}</p>'
                    html_content += "</div>"

                # Add footer
                html_content += f'<div class="footer">Report Prepared by {meta["author"]}</div>'
                html_content += "</body></html>"

                with open(filename, 'w') as f:
                    f.write(html_content)
                    
                return os.path.abspath(filename)
            except Exception as e:
                messagebox.showerror("HTML Error", f"HTML generation failed: {str(e)}")
                raise

    def get_severity_color(self, severity, excel=False):
        color_map = {
            'Low': COLORS['success'],
            'Medium': COLORS['warning'],
            'High': COLORS['error'],
            'Critical': COLORS['error']
        }
        return color_map.get(severity, COLORS['text'])

    def word_severity_color(self, severity):
        # Convert to title case for consistency
        severity = severity.title()
        color_map = {
            'Critical': COLORS['error'],
            'High': COLORS['error'],
            'Medium': COLORS['warning'],
            'Low': COLORS['success']
        }
        return RGBColor.from_string(color_map.get(severity, COLORS['text'])[1:])

if __name__ == "__main__":
    root = tk.Tk()
    app = ReportGeneratorApp(root)
    root.mainloop()
